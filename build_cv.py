"""Gera o CV em portugues, ingles e espanhol (DOCX e PDF).

Um unico conteudo por idioma, mesmo layout. Coluna unica, sem tabelas,
icones, imagens, cabecalho ou rodape, para que os parsers de ATS (sistemas
de RH) leiam o texto na ordem correta. Fontes padrao (Helvetica no PDF,
Calibri no DOCX), titulos de secao com nomes reconhecidos pelos ATS e datas
no mesmo formato em todo o documento. A hierarquia visual vem de tamanho,
peso, cor e filetes, na mesma paleta cobalto do site.

Saidas (em public/, servidas pelo site):
  cv.pdf, cv.docx          (portugues)
  cv_en.pdf, cv_en.docx    (ingles)
  cv_es.pdf, cv_es.docx    (espanhol)

Metricas do GitHub (org quero-edu, 22/10/2025 a 22/10/2026, API de busca):
  634 PRs abertos, 540 mergeados, 26 repositorios,
  295 PRs de outras pessoas revisados (reviewed-by sem os proprios PRs).

Uso: .tools-venv/bin/python build_cv.py
"""

from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate, Spacer

# Paleta do site (DESIGN.md): cobalto escuro para impressao, neutros frios.
INK = "#14161A"
ACCENT = "#2141CC"
MUTED = "#545C6A"
RULE = "#D5D9E2"

NAME = "Ali Mohamed Khodr"

# Dados que nao mudam entre idiomas.
EMAIL = "ali.m.khodr@gmail.com"
LINKEDIN = "linkedin.com/in/alimohamedkhodr"
GITHUB = "github.com/alimkhodr"
WEBSITE = "alimkhodr.netlify.app"
PHONE_LOCAL = "(12) 98859-4305"
PHONE_INTL = "+55 12 98859-4305"

CONTENT = {
    "pt": {
        "file": "cv",
        "title": "Software Engineer | Full-stack: Vue.js, Nuxt, TypeScript, Ruby on Rails, Node.js",
        "location": "São José dos Campos, SP",
        "phone": PHONE_LOCAL,
        "sections": {
            "summary": "Resumo Profissional",
            "skills": "Competências Técnicas",
            "experience": "Experiência Profissional",
            "education": "Formação Acadêmica",
            "languages": "Idiomas",
        },
        "tech_label": "Tecnologias",
        "summary": (
            "Software Engineer full-stack com experiência em levar produtos web do zero à produção. "
            "Tech lead de um gateway de pagamentos embarcável (Cartão, PIX e Boleto) e de uma área do aluno "
            "multi-tenant na Qeevo Group. Trabalho com Vue.js, Nuxt, TypeScript, Ruby on Rails e Node.js, "
            "além de Docker, AWS e CI/CD. Nos últimos 12 meses: 634 pull requests abertos (540 mergeados) e "
            "295 pull requests de outros desenvolvedores revisados, em 26 repositórios."
        ),
        "skills": [
            ("Linguagens", "TypeScript, JavaScript, Ruby, Python, SQL, C#, Visual Basic"),
            ("Frontend", "Vue.js, Nuxt, React, Pinia, Vuex, Tailwind CSS, HTML, CSS"),
            ("Backend e APIs", "Ruby on Rails, Node.js, REST APIs, GraphQL, OAuth"),
            ("Infraestrutura e DevOps", "Docker, AWS, CI/CD, GitHub Actions, Linux, Netlify, Vercel"),
            ("Dados", "PostgreSQL, MySQL, SQL Server, Supabase, modelagem relacional, pipelines ETL"),
            ("Domínios e práticas", "Pagamentos (Cartão, PIX, PIX recorrente, Boleto, reembolso), "
                                    "arquitetura multi-tenant, design systems, testes A/B, code review, "
                                    "testes automatizados, Git, metodologias ágeis, Claude Code"),
        ],
        "experience": [
            {
                "company": "Qeevo Group (ex-Quero Educação)",
                "location": "São José dos Campos, SP",
                "roles": [
                    {
                        "title": "Software Engineer",
                        "period": "2025 - atual",
                        "bullets": [
                            "Liderei tecnicamente a construção do zero de um gateway de pagamentos embarcável via "
                            "iframe para parceiros: Cartão, PIX e Boleto, cupons, temas por marca, CI/CD e Docker "
                            "(cerca de 33 pull requests entre backend e frontend).",
                            "Liderei a área do aluno multi-tenant, também criada do zero: histórico de pagamentos, "
                            "reembolso self-service, validação de chave PIX e localização por tenant (cerca de 24 "
                            "pull requests como owner técnico do produto).",
                            "Coordenei um fluxo de reembolso ponta a ponta em mais de 6 repositórios e vários times, "
                            "com reembolso via PIX, fallback automático por e-mail e painel de auditoria.",
                            "Habilitei o PIX recorrente como método de pagamento de assinaturas, do controller no "
                            "monolito Ruby ao modal de pagamento no frontend.",
                            "Revisei 295 pull requests de outros desenvolvedores nos últimos 12 meses, contribuí com "
                            "o design system compartilhado da organização e implementei testes A/B de cross-sell "
                            "pós-checkout.",
                        ],
                        "stack": "Vue.js, Nuxt, TypeScript, Ruby on Rails, Node.js, Docker, AWS, PostgreSQL, CI/CD",
                    },
                    {
                        "title": "Desenvolvedor Full-stack",
                        "period": "2024 - 2025",
                        "bullets": [
                            "Desenvolvi funcionalidades full-stack em um marketplace de educação com múltiplas "
                            "marcas e tenants.",
                            "Reformulei as páginas de pagamento de assinatura e implementei cross-sell multi-oferta "
                            "no pós-checkout, com estado controlado por cookies.",
                        ],
                        "stack": "Vue.js, React, Nuxt, Ruby, AWS, Docker",
                    },
                ],
            },
            {
                "company": "APTIV",
                "location": "São José dos Campos, SP",
                "roles": [
                    {
                        "title": "Desenvolvedor de Ferramentas Internas",
                        "period": "2022 - 2024",
                        "bullets": [
                            "Desenvolvi ferramentas internas e automações de processos industriais integradas a "
                            "bancos SQL Server em produção.",
                            "Apresentei à liderança um projeto interno de três meses, com reconhecimento formal.",
                        ],
                        "stack": "Visual Basic, C#, SQL, React, Node.js, Docker",
                    },
                ],
            },
        ],
        "education": [
            ("Análise e Desenvolvimento de Sistemas", "Universidade do Vale do Paraíba (Univap)",
             "2023 - 2025", "concluído"),
            ("Banco de Dados (Tecnólogo)", "Fatec São José dos Campos", "2022 - atual", ""),
            ("Técnico em Desenvolvimento de Sistemas", "ETEC Profª Ilza Nascimento Pintus", "2019 - 2021", ""),
        ],
        "languages": [
            ("Português", "nativo"),
            ("Inglês", "avançado (C1)"),
        ],
    },
    "en": {
        "file": "cv_en",
        "title": "Software Engineer | Full-stack: Vue.js, Nuxt, TypeScript, Ruby on Rails, Node.js",
        "location": "São José dos Campos, SP, Brazil",
        "phone": PHONE_INTL,
        "sections": {
            "summary": "Professional Summary",
            "skills": "Technical Skills",
            "experience": "Professional Experience",
            "education": "Education",
            "languages": "Languages",
        },
        "tech_label": "Technologies",
        "summary": (
            "Full-stack Software Engineer who takes web products from scratch to production. Tech lead of an "
            "embeddable payment gateway (Card, PIX and Boleto) and of a multi-tenant student portal at Qeevo "
            "Group. Works with Vue.js, Nuxt, TypeScript, Ruby on Rails and Node.js, plus Docker, AWS and CI/CD. "
            "Last 12 months: 634 pull requests opened (540 merged) and 295 pull requests from other developers "
            "reviewed, across 26 repositories."
        ),
        "skills": [
            ("Languages", "TypeScript, JavaScript, Ruby, Python, SQL, C#, Visual Basic"),
            ("Frontend", "Vue.js, Nuxt, React, Pinia, Vuex, Tailwind CSS, HTML, CSS"),
            ("Backend and APIs", "Ruby on Rails, Node.js, REST APIs, GraphQL, OAuth"),
            ("Infrastructure and DevOps", "Docker, AWS, CI/CD, GitHub Actions, Linux, Netlify, Vercel"),
            ("Data", "PostgreSQL, MySQL, SQL Server, Supabase, relational modeling, ETL pipelines"),
            ("Domains and practices", "Payments (Card, PIX, recurring PIX, Boleto, refunds), multi-tenant "
                                      "architecture, design systems, A/B testing, code review, automated "
                                      "testing, Git, Agile, Claude Code"),
        ],
        "experience": [
            {
                "company": "Qeevo Group (formerly Quero Educação)",
                "location": "São José dos Campos, SP, Brazil",
                "roles": [
                    {
                        "title": "Software Engineer",
                        "period": "2025 - present",
                        "bullets": [
                            "Led the technical build, from scratch, of an embeddable payment gateway for partners "
                            "(iframe): Card, PIX and Boleto flows, coupons, per-brand theming, CI/CD and Docker "
                            "(about 33 pull requests across backend and frontend).",
                            "Led a multi-tenant student portal, also built from scratch: payment history, "
                            "self-service refunds, PIX key validation and per-tenant localization (about 24 pull "
                            "requests as the product's technical owner).",
                            "Coordinated an end-to-end refund flow across more than 6 repositories and several "
                            "teams, with PIX refunds, automatic e-mail fallback and an audit panel.",
                            "Enabled recurring PIX as a subscription payment method, from the controller in the "
                            "Ruby monolith to the frontend payment modal.",
                            "Reviewed 295 pull requests from other developers in the last 12 months, contributed to "
                            "the organization's shared design system and implemented A/B tests for post-checkout "
                            "cross-sell.",
                        ],
                        "stack": "Vue.js, Nuxt, TypeScript, Ruby on Rails, Node.js, Docker, AWS, PostgreSQL, CI/CD",
                    },
                    {
                        "title": "Full-stack Developer",
                        "period": "2024 - 2025",
                        "bullets": [
                            "Built full-stack features for an education marketplace with multiple brands and tenants.",
                            "Redesigned the subscription payment pages and implemented multi-offer post-checkout "
                            "cross-sell with cookie-based state.",
                        ],
                        "stack": "Vue.js, React, Nuxt, Ruby, AWS, Docker",
                    },
                ],
            },
            {
                "company": "APTIV",
                "location": "São José dos Campos, SP, Brazil",
                "roles": [
                    {
                        "title": "Internal Tools Developer",
                        "period": "2022 - 2024",
                        "bullets": [
                            "Built internal tools and industrial process automations integrated with production "
                            "SQL Server databases.",
                            "Presented a three-month internal project to leadership and received formal recognition.",
                        ],
                        "stack": "Visual Basic, C#, SQL, React, Node.js, Docker",
                    },
                ],
            },
        ],
        "education": [
            ("Systems Analysis and Development", "Universidade do Vale do Paraíba (Univap)",
             "2023 - 2025", "completed"),
            ("Database Technology (Technologist degree)", "Fatec São José dos Campos", "2022 - present", ""),
            ("Systems Development Technician", "ETEC Profª Ilza Nascimento Pintus", "2019 - 2021", ""),
        ],
        "languages": [
            ("Portuguese", "native"),
            ("English", "advanced (C1)"),
        ],
    },
    "es": {
        "file": "cv_es",
        "title": "Ingeniero de Software | Full-stack: Vue.js, Nuxt, TypeScript, Ruby on Rails, Node.js",
        "location": "São José dos Campos, SP, Brasil",
        "phone": PHONE_INTL,
        "sections": {
            "summary": "Resumen Profesional",
            "skills": "Competencias Técnicas",
            "experience": "Experiencia Profesional",
            "education": "Formación Académica",
            "languages": "Idiomas",
        },
        "tech_label": "Tecnologías",
        "summary": (
            "Ingeniero de software full-stack que lleva productos web desde cero hasta producción. Tech lead de "
            "una pasarela de pagos embebible (Tarjeta, PIX y Boleto) y de un portal de estudiantes multi-tenant "
            "en Qeevo Group. Trabaja con Vue.js, Nuxt, TypeScript, Ruby on Rails y Node.js, además de Docker, "
            "AWS y CI/CD. Últimos 12 meses: 634 pull requests abiertos (540 integrados) y 295 pull requests de "
            "otros desarrolladores revisados, en 26 repositorios."
        ),
        "skills": [
            ("Lenguajes", "TypeScript, JavaScript, Ruby, Python, SQL, C#, Visual Basic"),
            ("Frontend", "Vue.js, Nuxt, React, Pinia, Vuex, Tailwind CSS, HTML, CSS"),
            ("Backend y APIs", "Ruby on Rails, Node.js, APIs REST, GraphQL, OAuth"),
            ("Infraestructura y DevOps", "Docker, AWS, CI/CD, GitHub Actions, Linux, Netlify, Vercel"),
            ("Datos", "PostgreSQL, MySQL, SQL Server, Supabase, modelado relacional, pipelines ETL"),
            ("Dominios y prácticas", "Pagos (Tarjeta, PIX, PIX recurrente, Boleto, reembolsos), arquitectura "
                                     "multi-tenant, design systems, pruebas A/B, code review, pruebas "
                                     "automatizadas, Git, metodologías ágiles, Claude Code"),
        ],
        "experience": [
            {
                "company": "Qeevo Group (antes Quero Educação)",
                "location": "São José dos Campos, SP, Brasil",
                "roles": [
                    {
                        "title": "Ingeniero de Software",
                        "period": "2025 - actual",
                        "bullets": [
                            "Lideré técnicamente la construcción desde cero de una pasarela de pagos embebible para "
                            "socios (iframe): flujos de Tarjeta, PIX y Boleto, cupones, temas por marca, CI/CD y "
                            "Docker (unos 33 pull requests entre backend y frontend).",
                            "Lideré un portal de estudiantes multi-tenant, también construido desde cero: historial "
                            "de pagos, reembolsos de autoservicio, validación de clave PIX y localización por tenant "
                            "(unos 24 pull requests como responsable técnico del producto).",
                            "Coordiné un flujo de reembolso de punta a punta en más de 6 repositorios y varios "
                            "equipos, con reembolso por PIX, respaldo automático por correo electrónico y panel de "
                            "auditoría.",
                            "Habilité el PIX recurrente como método de pago de suscripciones, desde el controller "
                            "del monolito Ruby hasta el modal de pago del frontend.",
                            "Revisé 295 pull requests de otros desarrolladores en los últimos 12 meses, contribuí al "
                            "design system compartido de la organización e implementé pruebas A/B de cross-sell "
                            "post-checkout.",
                        ],
                        "stack": "Vue.js, Nuxt, TypeScript, Ruby on Rails, Node.js, Docker, AWS, PostgreSQL, CI/CD",
                    },
                    {
                        "title": "Desarrollador Full-stack",
                        "period": "2024 - 2025",
                        "bullets": [
                            "Desarrollé funcionalidades full-stack en un marketplace educativo con múltiples marcas "
                            "y tenants.",
                            "Rediseñé las páginas de pago de suscripciones e implementé cross-sell multi-oferta "
                            "post-checkout, con estado controlado por cookies.",
                        ],
                        "stack": "Vue.js, React, Nuxt, Ruby, AWS, Docker",
                    },
                ],
            },
            {
                "company": "APTIV",
                "location": "São José dos Campos, SP, Brasil",
                "roles": [
                    {
                        "title": "Desarrollador de Herramientas Internas",
                        "period": "2022 - 2024",
                        "bullets": [
                            "Desarrollé herramientas internas y automatizaciones de procesos industriales integradas "
                            "con bases SQL Server en producción.",
                            "Presenté a la dirección un proyecto interno de tres meses, con reconocimiento formal.",
                        ],
                        "stack": "Visual Basic, C#, SQL, React, Node.js, Docker",
                    },
                ],
            },
        ],
        "education": [
            ("Análisis y Desarrollo de Sistemas", "Universidade do Vale do Paraíba (Univap)",
             "2023 - 2025", "completado"),
            ("Tecnología en Base de Datos", "Fatec São José dos Campos", "2022 - actual", ""),
            ("Técnico en Desarrollo de Sistemas", "ETEC Profª Ilza Nascimento Pintus", "2019 - 2021", ""),
        ],
        "languages": [
            ("Portugués", "nativo"),
            ("Inglés", "avanzado (C1)"),
        ],
    },
}

SEPARATOR = "  |  "


def _contact_lines(lang):
    """Duas linhas de contato, quebradas de proposito para nao cortar um link no meio."""
    c = CONTENT[lang]
    first = [c["location"], EMAIL, c["phone"]]
    second = [LINKEDIN, GITHUB, WEBSITE]
    return [SEPARATOR.join(first), SEPARATOR.join(second)]


def _education_line(entry):
    course, school, period, status = entry
    meta = period + (f", {status}" if status else "")
    return course, f"{school}{SEPARATOR}{meta}"


def _languages_line(lang):
    return SEPARATOR.join(f"{name}: {level}" for name, level in CONTENT[lang]["languages"])


def _keywords(lang, limit=255):
    """Palavras-chave para os metadados do arquivo: as tecnologias da secao de competencias.

    O DOCX aceita no maximo 255 caracteres nesse campo, entao a lista para antes do limite."""
    items = []
    for _, values in CONTENT[lang]["skills"][:5]:
        items += [v.strip() for v in values.split(",")]
    text = ""
    for item in dict.fromkeys(items):
        candidate = f"{text}, {item}" if text else item
        if len(candidate) > limit:
            break
        text = candidate
    return text


# ---------------------------------------------------------------- DOCX

def _rgb(hex_color):
    return RGBColor.from_string(hex_color.lstrip("#"))


def _bottom_border(paragraph, hex_color, size=8):
    """Filete embaixo do paragrafo (usado nos titulos de secao)."""
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), hex_color.lstrip("#"))
    borders.append(bottom)
    p_pr.append(borders)


def build_docx(lang, path):
    c = CONTENT[lang]
    doc = Document()
    for section in doc.sections:
        section.page_height, section.page_width = Cm(29.7), Cm(21)
        section.top_margin = section.bottom_margin = Cm(1.3)
        section.left_margin = section.right_margin = Cm(1.7)

    props = doc.core_properties
    props.title = f"{NAME} - CV"
    props.author = NAME
    props.subject = c["title"]
    props.keywords = _keywords(lang)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    normal.font.color.rgb = _rgb(INK)
    normal.paragraph_format.space_after = Pt(2)
    normal.paragraph_format.line_spacing = 1.1

    bullet_style = doc.styles["List Bullet"]
    bullet_style.font.size = Pt(10)
    bullet_style.paragraph_format.space_after = Pt(1)

    name_p = doc.add_paragraph()
    name_p.paragraph_format.space_after = Pt(0)
    run = name_p.add_run(NAME)
    run.bold = True
    run.font.size = Pt(22)

    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_after = Pt(2)
    t_run = title_p.add_run(c["title"])
    t_run.bold = True
    t_run.font.size = Pt(10.5)
    t_run.font.color.rgb = _rgb(ACCENT)

    lines = _contact_lines(lang)
    contact_p = doc.add_paragraph()
    contact_p.paragraph_format.space_after = Pt(4)
    for i, line in enumerate(lines):
        c_run = contact_p.add_run(line)
        c_run.font.size = Pt(9)
        c_run.font.color.rgb = _rgb(MUTED)
        if i < len(lines) - 1:
            c_run.add_break()
    _bottom_border(contact_p, ACCENT, size=12)

    def heading(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(9)
        h.paragraph_format.space_after = Pt(4)
        h.paragraph_format.keep_with_next = True
        r = h.add_run(text.upper())
        r.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = _rgb(ACCENT)
        _bottom_border(h, RULE, size=6)

    def muted(paragraph, text, size=9):
        r = paragraph.add_run(text)
        r.font.size = Pt(size)
        r.font.color.rgb = _rgb(MUTED)
        return r

    heading(c["sections"]["summary"])
    doc.add_paragraph(c["summary"])

    heading(c["sections"]["skills"])
    for label, items in c["skills"]:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        p.add_run(f"{label}: ").bold = True
        p.add_run(items)

    heading(c["sections"]["experience"])
    for company in c["experience"]:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(5)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.keep_with_next = True
        co = p.add_run(company["company"])
        co.bold = True
        co.font.size = Pt(10.5)
        muted(p, f"{SEPARATOR}{company['location']}")
        for role in company["roles"]:
            r = doc.add_paragraph()
            r.paragraph_format.space_before = Pt(3)
            r.paragraph_format.space_after = Pt(1)
            r.paragraph_format.keep_with_next = True
            t = r.add_run(role["title"])
            t.bold = True
            t.font.color.rgb = _rgb(ACCENT)
            muted(r, f"{SEPARATOR}{role['period']}")
            for b in role["bullets"]:
                doc.add_paragraph(b, style="List Bullet")
            s = doc.add_paragraph()
            s.paragraph_format.space_after = Pt(3)
            muted(s, f"{c['tech_label']}: {role['stack']}")

    heading(c["sections"]["education"])
    for entry in c["education"]:
        course, meta = _education_line(entry)
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.add_run(course).bold = True
        muted(p, f"{SEPARATOR}{meta}")

    heading(c["sections"]["languages"])
    doc.add_paragraph(_languages_line(lang))

    doc.save(path)


# ---------------------------------------------------------------- PDF

def build_pdf(lang, path):
    c = CONTENT[lang]
    base = ParagraphStyle("base", fontName="Helvetica", fontSize=9.2, leading=12.3,
                          spaceAfter=1.5, textColor=colors.HexColor(INK))
    name = ParagraphStyle("name", parent=base, fontName="Helvetica-Bold", fontSize=22,
                          leading=25, spaceAfter=1)
    title = ParagraphStyle("title", parent=base, fontName="Helvetica-Bold", fontSize=10.2,
                           leading=13, textColor=colors.HexColor(ACCENT), spaceAfter=3)
    contact = ParagraphStyle("contact", parent=base, fontSize=8.6, leading=11.5,
                             textColor=colors.HexColor(MUTED), spaceAfter=0)
    head = ParagraphStyle("head", parent=base, fontName="Helvetica-Bold", fontSize=9.4,
                          leading=12, textColor=colors.HexColor(ACCENT), spaceBefore=9, spaceAfter=3)
    company = ParagraphStyle("company", parent=base, fontName="Helvetica-Bold", fontSize=10.2,
                             leading=13, spaceBefore=4, spaceAfter=0)
    role = ParagraphStyle("role", parent=base, fontName="Helvetica-Bold", fontSize=9.4,
                          leading=12.3, textColor=colors.HexColor(ACCENT), spaceBefore=2.5, spaceAfter=1.5)
    stack = ParagraphStyle("stack", parent=base, fontSize=8.5, leading=11.2,
                           textColor=colors.HexColor(MUTED), spaceBefore=0.5, spaceAfter=3)
    bullet = ParagraphStyle("bullet", parent=base, leftIndent=11, bulletIndent=1.5, spaceAfter=1.5)
    meta = f"<font name='Helvetica' size=8.5 color='{MUTED}'>"

    def esc(t):
        return t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    def section(text):
        # Titulo de secao com filete abaixo, desenhado pelo flowable HRFlowable.
        return [Paragraph(text.upper(), head),
                HRFlowable(width="100%", thickness=0.6, color=colors.HexColor(RULE), spaceBefore=0, spaceAfter=4)]

    story = [
        Paragraph(NAME, name),
        Paragraph(esc(c["title"]), title),
        Paragraph("<br/>".join(esc(line) for line in _contact_lines(lang)), contact),
        HRFlowable(width="100%", thickness=1.4, color=colors.HexColor(ACCENT), spaceBefore=5, spaceAfter=0),
    ]

    story += section(c["sections"]["summary"])
    story.append(Paragraph(esc(c["summary"]), base))

    story += section(c["sections"]["skills"])
    for label, items in c["skills"]:
        story.append(Paragraph(f"<b>{esc(label)}:</b> {esc(items)}", base))

    story += section(c["sections"]["experience"])
    for co in c["experience"]:
        story.append(Paragraph(f"{esc(co['company'])}{meta}{SEPARATOR}{esc(co['location'])}</font>", company))
        for ro in co["roles"]:
            story.append(Paragraph(f"{esc(ro['title'])}{meta}{SEPARATOR}{esc(ro['period'])}</font>", role))
            for b in ro["bullets"]:
                story.append(Paragraph(esc(b), bullet, bulletText="•"))
            story.append(Paragraph(f"{esc(c['tech_label'])}: {esc(ro['stack'])}", stack))

    story += section(c["sections"]["education"])
    for entry in c["education"]:
        course, info = _education_line(entry)
        story.append(Paragraph(f"<b>{esc(course)}</b>{meta}{SEPARATOR}{esc(info)}</font>", base))
        story.append(Spacer(1, 1.5))

    story += section(c["sections"]["languages"])
    story.append(Paragraph(esc(_languages_line(lang)), base))

    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=1.7 * cm, rightMargin=1.7 * cm,
                            topMargin=1.3 * cm, bottomMargin=1.3 * cm,
                            title=f"{NAME} - CV", author=NAME, subject=c["title"],
                            keywords=_keywords(lang), lang={"pt": "pt-BR", "en": "en", "es": "es"}[lang])
    doc.build(story)


OUTPUT_DIR = Path(__file__).parent / "public"


if __name__ == "__main__":
    OUTPUT_DIR.mkdir(exist_ok=True)
    for code, content in CONTENT.items():
        build_docx(code, str(OUTPUT_DIR / f"{content['file']}.docx"))
        build_pdf(code, str(OUTPUT_DIR / f"{content['file']}.pdf"))
        print(f"public/{content['file']}.pdf e public/{content['file']}.docx gerados")
